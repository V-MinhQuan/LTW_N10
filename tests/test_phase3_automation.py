import os
import sys
import json
import io

import django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'QLVB.settings')
sys.path.insert(0, 'QLVB')
django.setup()

from django.test import RequestFactory
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from docx import Document

from web.models import VanBanDen, VanBanDi, UserAccount, DonViBenTrong, DonViBenNgoai
from web.views import (
    api_ocr_extract_document,
    api_ai_summarize_document,
    api_ai_suggest_assignment
)
from web.ai_services import (
    extract_text_from_file,
    parse_ocr_administrative_metadata,
    classify_and_suggest_assignment,
    generate_smart_summary
)

factory = RequestFactory()
user = UserAccount.objects.first()
if not user:
    user = UserAccount.objects.create(username='ai_test_user', HoTen='Nguyen Van AI')

print("=" * 70)
print("BẮT ĐẦU KIỂM THỬ GIAI ĐOẠN 3 (PHASE 3 AUTOMATION TESTS)")
print("=" * 70)

# ==============================================================================
# TEST 1: AUTO-01 - OCR & TRÍCH XUẤT THÔNG TIN HÀNH CHÍNH TỰ ĐỘNG
# ==============================================================================
print("\n>>> [TEST 1] AUTO-01: OCR & Trích xuất thông tin hành chính từ tệp")

# 1.1 Tạo tệp DOCX mẫu chuẩn thể thức hành chính Việt Nam
doc = Document()
doc.add_paragraph("ỦY BAN NHÂN DÂN THÀNH PHỐ HÀ NỘI")
doc.add_paragraph("SỞ DU LỊCH")
doc.add_paragraph("Số: 99/TB-SDL")
doc.add_paragraph("Hà Nội, ngày 28 tháng 10 năm 2026")
doc.add_paragraph("THÔNG BÁO")
doc.add_paragraph("V/v tổ chức hội nghị xúc tiến điểm đến và phát triển tour du lịch mùa đông 2026")
doc.add_paragraph("Kính gửi: Công ty Du lịch Vietasia Travel.")
doc.add_paragraph("1. Triển khai kế hoạch quảng bá sản phẩm lữ hành cao cấp quý IV/2026.")
doc.add_paragraph("2. Đề nghị các doanh nghiệp hoàn tất gửi danh sách đại biểu trước ngày 05/11/2026.")
doc.add_paragraph("GIÁM ĐỐC")
doc.add_paragraph("Đặng Hương Giang")

doc_stream = io.BytesIO()
doc.save(doc_stream)
doc_stream.seek(0)
doc_bytes = doc_stream.getvalue()

uploaded_file = SimpleUploadedFile(
    "Thong_bao_hop_chuan_2026.docx",
    doc_bytes,
    content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

# Gửi yêu cầu OCR
req_ocr = factory.post(
    '/api/ai/ocr-extract/',
    {'tep_dinh_kem': uploaded_file}
)
req_ocr.user = user

resp_ocr = api_ocr_extract_document(req_ocr)
assert resp_ocr.status_code == 200, f"Expected 200, got {resp_ocr.status_code}"
ocr_data = json.loads(resp_ocr.content.decode('utf-8'))
assert ocr_data['status'] == 'success'
d = ocr_data['data']

print("  Kết quả OCR nhận diện:")
print("   - Số ký hiệu:", d['so_ky_hieu'])
print("   - Loại văn bản:", d['loai_van_ban'])
print("   - Ngày ban hành:", d['ngay_ban_hanh'])
print("   - Cơ quan ban hành:", d['co_quan_ban_hanh'])
print("   - Trích yếu:", d['trich_yeu'])
print("   - Người ký:", d['nguoi_ky'])

assert '99/TB-SDL' in d['so_ky_hieu'], f"Expected 99/TB-SDL, got {d['so_ky_hieu']}"
assert d['loai_van_ban'] == 'Thông báo'
assert d['ngay_ban_hanh'] == '2026-10-28'
assert 'SỞ DU LỊCH' in d['co_quan_ban_hanh'] or 'ỦY BAN' in d['co_quan_ban_hanh']
assert 'hội nghị xúc tiến' in d['trich_yeu'].lower()
assert 'Đặng Hương Giang' in d['nguoi_ky']

print("  [PASS] 1.1 Tự động trích xuất đầy đủ 6/6 trường metadata từ tệp chuẩn xác 100%!")


# ==============================================================================
# TEST 2: AUTO-07 - PHÂN LOẠI VĂN BẢN & GỢI Ý PHÂN CÔNG CHO LÃNH ĐẠO
# ==============================================================================
print("\n>>> [TEST 2] AUTO-07: Phân loại văn bản & Gợi ý phân công cho Lãnh đạo")

# 2.1 Kiểm tra phân loại lĩnh vực
test_cases = [
    ("Thanh toán tiền tạm ứng và quyết toán thuế quý 3", "Tài chính - Kế toán"),
    ("Lịch trình đoàn khách tham quan Hạ Long và phân công hướng dẫn viên tour", "Kinh doanh Lữ hành & Điều hành Tour"),
    ("Thông báo tuyển dụng chuyên viên điều hành tour và chính sách nghỉ phép", "Hành chính - Nhân sự"),
    ("Giải quyết tranh chấp hợp đồng kinh tế và rà soát thủ tục pháp lý", "Pháp chế & Hợp tác đối tác"),
]

for text_sample, expected_domain in test_cases:
    cls_res = classify_and_suggest_assignment(text_sample, text_sample, "Công văn")
    print(f"  Văn bản: '{text_sample[:45]}...' -> Phân loại: {cls_res['linh_vuc']}")
    assert cls_res['linh_vuc'] == expected_domain, f"Expected {expected_domain}, got {cls_res['linh_vuc']}"
    assert cls_res['user_name'] != '', "User name should not be empty"

print("  [PASS] 2.1 Phân loại lĩnh vực và gợi ý người phụ trách hoạt động chuẩn xác!")

# 2.2 Kiểm tra API gợi ý phân công
req_sug = factory.get('/api/ai/suggest-assignment/', {
    'trich_yeu': 'Kế hoạch tổ chức đoàn khách du lịch Châu Âu mùa thu 2026',
    'loai_vb': 'Kế hoạch'
})
req_sug.user = user
resp_sug = api_ai_suggest_assignment(req_sug)
assert resp_sug.status_code == 200
sug_json = json.loads(resp_sug.content.decode('utf-8'))
assert sug_json['status'] == 'success'
assert sug_json['data']['linh_vuc'] == 'Kinh doanh Lữ hành & Điều hành Tour'
print(f"  [PASS] 2.2 API /api/ai/suggest-assignment/ trả về gợi ý: {sug_json['data']['user_name']} ({sug_json['data']['phong_ban_ten']})")


# ==============================================================================
# TEST 3: AUTO-08 - TÓM TẮT THÔNG MINH BẰNG AI (EXECUTIVE SMART SUMMARY)
# ==============================================================================
print("\n>>> [TEST 3] AUTO-08: Nghiên cứu & Tóm tắt thông minh bằng AI")

# 3.1 Tạo văn bản đến thử nghiệm
vbd_test = VanBanDen.objects.create(
    SoKyHieu='123/TB-TEST',
    TrichYeu='V/v hướng dẫn triển khai công tác an toàn phòng chống cháy nổ tại các chi nhánh du lịch',
    LoaiVanBan='Thông báo',
    NgayBanHanh=timezone.now(),
    TrangThai='DANG_XU_LY'
)

# 3.2 Gọi API tóm tắt AI
req_sum = factory.post(f'/api/ai/{vbd_test.VanBanDenID}/summarize/?type=den')
req_sum.user = user
resp_sum = api_ai_summarize_document(req_sum, pk=vbd_test.VanBanDenID)
assert resp_sum.status_code == 200
sum_json = json.loads(resp_sum.content.decode('utf-8'))
assert sum_json['status'] == 'success'
assert 'summary_html' in sum_json
assert 'summary_data' in sum_json

sum_d = sum_json['summary_data']
print("  Nội dung tóm tắt thông minh:")
print("   - Trọng tâm:", sum_d['muc_dich_cot_loi'])
print("   - Số ý chính:", len(sum_d['diem_chinh']))
print("   - Điểm chính đầu tiên:", sum_d['diem_chinh'][0])
print("   - Thời hạn:", sum_d['thoi_han_yeu_cau'])
print("   - Đơn vị gợi ý:", sum_d['don_vi_chu_tri_goi_y'])

assert 3 <= len(sum_d['diem_chinh']) <= 5
assert len(sum_d['muc_dich_cot_loi']) > 10

# 3.3 Kiểm tra tóm tắt đã được lưu cache trong CSDL
vbd_test.refresh_from_db()
assert vbd_test.TomTatAI is not None and len(vbd_test.TomTatAI) > 50
assert vbd_test.LinhVucAI is not None
print("  [PASS] 3.1 Tóm tắt AI được sinh và lưu cache vào CSDL thành công!")

# 3.4 Gọi lại lần 2 -> Trả về kết quả từ cache (cached: True)
req_sum_cache = factory.get(f'/api/ai/{vbd_test.VanBanDenID}/summarize/?type=den')
req_sum_cache.user = user
resp_sum_cache = api_ai_summarize_document(req_sum_cache, pk=vbd_test.VanBanDenID)
assert resp_sum_cache.status_code == 200
sum_cache_json = json.loads(resp_sum_cache.content.decode('utf-8'))
assert sum_cache_json['cached'] is True
print("  [PASS] 3.2 Cache tóm tắt AI hoạt động hoàn hảo, tải tức thời không tốn tài nguyên!")

# Cleanup
vbd_test.delete()

print("\n" + "=" * 70)
print("TẤT CẢ CÁC BÀI TEST GIAI ĐOẠN 3 ĐÃ ĐẠT 100% THÀNH CÔNG!")
print("=" * 70)

